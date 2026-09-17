"""The reader streams: the head, the count and the batches come off one pass
each and agree with the whole table, for CSV and XLSX alike."""

from __future__ import annotations

import io

from lge.ingest.batches import plan_batches, slice_rows
from lge.ingest.reader import count_rows, iter_batches, iter_rows, read_head, read_table

N = 1203


def _csv(tmp_path, n=N):
    p = tmp_path / "big.csv"
    lines = ["Firm,Site,City"]
    for i in range(n):
        lines.append(f"Firm {i},firm{i}.com,Atlanta")
        if i % 400 == 0:
            lines.append("")  # blank lines are not rows
    p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return p


def _xlsx(tmp_path, n=120):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Firm", "Site"])
    for i in range(n):
        ws.append([f"Firm {i}", f"firm{i}.com"])
    buf = io.BytesIO()
    wb.save(buf)
    p = tmp_path / "book.xlsx"
    p.write_bytes(buf.getvalue())
    return p


def test_head_count_and_batches_agree_with_the_whole_table(tmp_path):
    p = _csv(tmp_path)
    whole = read_table(p)
    assert whole.row_count == N and count_rows(p) == N
    head = read_head(p, 5)
    assert head.headers == whole.headers and head.rows == whole.rows[:5]
    assert read_head(p, 0).headers == whole.headers and read_head(p, 0).rows == []

    batches = plan_batches(N, 500)
    seen = list(iter_batches(p, batches))
    assert [b.number for b, _ in seen] == [1, 2, 3]
    assert [len(rows) for _, rows in seen] == [500, 500, 203]
    for b, rows in seen:
        assert rows == slice_rows(whole.rows, b)
    assert sum(1 for _ in iter_rows(p)) == N


def test_only_the_planned_batches_are_read_and_a_gap_is_skipped(tmp_path):
    p = _csv(tmp_path)
    whole = read_table(p)
    batches = plan_batches(N, 500)
    seen = list(iter_batches(p, [batches[2], batches[0]]))  # out of order, batch 2 missing
    assert [b.number for b, _ in seen] == [1, 3]
    assert seen[1][1] == slice_rows(whole.rows, batches[2])
    assert list(iter_batches(p, [])) == []


def test_a_plan_past_the_end_of_the_sheet_yields_what_exists(tmp_path):
    p = _csv(tmp_path, n=10)
    batches = plan_batches(60, 50)  # planned for more rows than the file has
    seen = list(iter_batches(p, batches))
    assert len(seen) == 1 and seen[0][0].number == 1 and len(seen[0][1]) == 10


def test_xlsx_streams_the_same_way(tmp_path):
    p = _xlsx(tmp_path)
    whole = read_table(p)
    assert whole.format == "xlsx" and whole.row_count == 120 and count_rows(p) == 120
    assert read_head(p, 3).rows == whole.rows[:3]
    seen = list(iter_batches(p, plan_batches(120, 50)))
    assert [len(r) for _, r in seen] == [50, 50, 20]
    assert seen[2][1] == whole.rows[100:]
